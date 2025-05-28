from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx2pdf import convert
import os


def create_document(
    logo_path,
    submit_by_list,
    submit_to,
    subject,
    roll_no_list,
    logo_width_inch=5.0,
    font_size_pt=20,
    output_path="submission.docx",
):
    doc = Document()

    for i in range(len(submit_by_list)):
        if i > 0:
            doc.add_page_break()

        if os.path.exists(logo_path):
            run = doc.add_paragraph().add_run()
            run.add_picture(logo_path, width=Inches(logo_width_inch))
            doc.paragraphs[-1].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        else:
            print("Logo not found at:", logo_path)

        doc.add_paragraph()  # Spacer

        def add_centered_kv_line(key, value):
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            run_key = p.add_run(f"{key}: ")
            run_key.bold = True
            run_key.font.size = Pt(font_size_pt)
            run_key.font.name = "Times New Roman"

            run_val = p.add_run(str(value))
            run_val.font.size = Pt(font_size_pt)
            run_val.font.name = "Times New Roman"

        add_centered_kv_line("Submitted by", submit_by_list[i])
        add_centered_kv_line("Submitted to", submit_to)
        add_centered_kv_line(
            "Subject", subject[0] if isinstance(subject, list) else subject
        )
        add_centered_kv_line("Roll no.", roll_no_list[i])

    doc.save(output_path)
    print(f"Document saved to: {output_path}")

    # Convert to PDF
    convert(output_path, output_path.replace(".docx", ".pdf"))
    print("PDF saved to:", output_path.replace(".docx", ".pdf"))


# ====== Example Call with Your Data ======

create_document(
    logo_path=r"C:\Users\punit\Desktop\logo.png",
    submit_by_list=["Punit Soni", "Karan Lodhi", "Kartik Chaturvedi"],
    # submit_by_list=["Kartik Chaturvedi"],
    submit_to="Mrs. Bhawna Nigam Ma'am",
    subject=["Data Analysis "],
    roll_no_list=["23I6093", "22I6035", "22I6036"],
    # roll_no_list=["22I6036"],
    logo_width_inch=4.0,
    font_size_pt=18,
    output_path="submission.docx",
)
